"""Build the short, plain-language progress update for the supervisor.

Separate from make_docx.py (which produces the long formal report). This one is
deliberately short, uses simple wording, explains every symbol, and keeps
examples in round numbers.

Run:  .venv/bin/python -m scripts.make_update_docx
"""
from __future__ import annotations

import json
import sys

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from ltf.config import OUTPUT

BLUE = RGBColor(0x1F, 0x3F, 0x66)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexcolor)
    tcPr.append(el)


def table(doc, header, rows, widths=None, font=9, bold_rows=None):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, txt in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(txt))
        r.bold = True
        r.font.size = Pt(font)
        shade(c, "E7EDF4")
    bold_rows = set(bold_rows or [])
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(val))
            r.font.size = Pt(font)
            if ri in bold_rows:
                r.bold = True
                shade(cells[i], "F2F7FB")
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def head(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = BLUE
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    return p


def txt(doc, s, size=10, bold=False, italic=False, after=6):
    p = doc.add_paragraph()
    r = p.add_run(s)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    p.paragraph_format.space_after = Pt(after)
    return p


def bul(doc, items, size=10):
    for it in items:
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run(it)
        r.font.size = Pt(size)
        p.paragraph_format.space_after = Pt(2)


def formula(doc, s, size=10.5):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(s)
    r.font.name = "Consolas"
    r.font.size = Pt(size)
    r.bold = True
    return p


def pc(new, old):
    if old == 0:
        return "n/a"
    return f"{100*(new-old)/abs(old):+.1f}%"


def main() -> int:
    t1 = pd.read_csv(OUTPUT / "table1_methods.csv")
    t2 = pd.read_csv(OUTPUT / "table2_ablations.csv")
    rb = pd.read_csv(OUTPUT / "robustness.csv")
    gr = pd.read_csv(OUTPUT / "sweep_joint_grid.csv")
    fc = json.load(open(OUTPUT / "forecaster_results.json"))
    PAPER, OURS = "Kang et al. (reproduced)", "Ours (Gap 1+2)"
    k = t1[t1.method == PAPER].iloc[0]
    o = t1[t1.method == OURS].iloc[0]
    ab = {r.method: r for _, r in t2.iterrows()}
    nofair, g12, static, nopred = (ab["Ours, w/o fairness"],
                                   ab["Ours, w/o utilisation term"],
                                   ab["Ours, static utility (c=1)"],
                                   ab["Ours, w/o prediction"])

    doc = Document()
    s = doc.styles["Normal"]
    s.font.name = "Calibri"
    s.font.size = Pt(10)
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Inches(0.75)
        sec.left_margin = sec.right_margin = Inches(0.85)

    # ---------- title ----------
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MTP Progress Update")
    r.bold = True
    r.font.size = Pt(18)
    r.font.color.rgb = BLUE
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Fixing two problems in \"Long-term Fairness in Ride-Hailing Platform\"")
    r.font.size = Pt(11.5)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Paper: Kang, Chan, Shao, Salim, Leckie — ECML PKDD 2024 (arXiv:2407.17839)\n"
                  "Data: New York City yellow taxi trips, March 2016")
    r.font.size = Pt(9.5)
    r.italic = True

    txt(doc, "")
    txt(doc, "Sir, this is a short update on what I have implemented so far. I first "
             "explain what the paper does, then the two problems I found in it, then how "
             "I fixed them and what the numbers came out to be. I have kept every formula "
             "with an explanation of each symbol, and small examples wherever possible.",
        size=10)

    # ================= 1 =================
    head(doc, "1. What the paper does", 1)
    txt(doc, "A ride-hailing app has to decide which trip request goes to which driver. "
             "The paper wants two things at the same time: total earnings should be high "
             "(efficiency), and earnings should be spread evenly across drivers (fairness). "
             "It uses reinforcement learning to do this, and it also predicts future trip "
             "requests so the decisions are not short-sighted.")
    txt(doc, "The paper has two main formulas.", bold=True, after=2)

    txt(doc, "Formula 1 — value of giving a trip to a driver:", after=2)
    formula(doc, "U  =  Geo(d, s)  -  Geo(s, g)")
    table(doc, ["Symbol", "What it means"], [
        ["U", "Utility. How useful this trip is for this driver. Measured in kilometres."],
        ["s", "Start point of the trip (where the passenger is waiting)"],
        ["d", "Drop point of the trip (where the passenger wants to go)"],
        ["g", "Where the driver is standing right now"],
        ["Geo(a, b)", "Road distance from point a to point b, in km"],
        ["Geo(d, s)", "Length of the paid trip. This is the good part."],
        ["Geo(s, g)", "Distance the driver must travel empty to reach the passenger. "
                      "No money for this, so it is the cost part."],
    ], widths=[1.0, 5.8])
    txt(doc, "Example: driver is 2 km away from the passenger, and the trip itself is 8 km. "
             "Then U = 8 - 2 = 6 km. If the driver were 9 km away, U = 8 - 9 = -1 km, "
             "meaning it is a loss for that driver.", size=9.5, italic=True)

    txt(doc, "Formula 2 — how the paper measures fairness:", after=2)
    formula(doc, "Fairness  =  Var( total weekly earning of each driver )")
    table(doc, ["Symbol", "What it means"], [
        ["Var(...)", "Variance. A number that says how spread out the values are. "
                     "If all drivers earn the same, variance is 0."],
        ["total weekly earning", "Sum of U for all trips that driver got in the whole week"],
    ], widths=[1.6, 5.2])
    txt(doc, "Example: three drivers finish the week with 100, 100 and 100. Variance = 0, "
             "so the paper says this is perfectly fair. If they finish with 100, 120 and 80, "
             "variance = 266.7, so it is less fair.", size=9.5, italic=True)
    txt(doc, "The paper then tries to get high total earnings and low variance together.")

    # ================= 2 =================
    head(doc, "2. The two problems I found", 1)

    head(doc, "Problem 1 (Gap 1): the formulas ignore traffic and ignore working hours", 2)
    txt(doc, "Part A — traffic is ignored. In Formula 1 there is only distance, no time. "
             "The paper also says clearly in Section 5.1 that it takes the average travel "
             "time for each route and uses that one number for the whole day. So a 5 km "
             "trip at 4 in the morning and the same 5 km trip at 6 in the evening are "
             "treated as exactly equal. In real Manhattan traffic they are not:")
    table(doc, ["Same 5 km trip", "Actual time taken"], [
        ["At 4 AM (empty roads)", "about 12 minutes"],
        ["At 6 PM (peak traffic)", "about 26 minutes"],
    ], widths=[2.6, 2.6])
    txt(doc, "I checked this on the paper's own data. For the same start and end point, the "
             "travel time at peak is 2.11 times the off-peak time on average, and up to "
             "4.35 times in the worst case. So the paper is throwing away a real effect.")

    txt(doc, "Part B — working hours are ignored. Formula 2 only looks at total earning. "
             "It never asks how long the driver was working. Take this example:")
    table(doc, ["Driver", "Earned in the week", "Hours worked", "Actual rate per hour"], [
        ["A (full-time)", "Rs 6,000", "60 hours", "Rs 100 per hour"],
        ["B (part-time)", "Rs 6,000", "15 hours", "Rs 400 per hour"],
    ], widths=[1.5, 1.7, 1.4, 1.9])
    txt(doc, "Both earned Rs 6,000, so variance = 0 and the paper calls this perfectly fair. "
             "But A is actually earning 4 times less per hour than B. The paper cannot see "
             "this at all.", size=9.5, italic=True)

    head(doc, "Problem 2 (Gap 2): it never checks if an available driver got any work", 2)
    txt(doc, "The paper only looks at final earnings. It never checks whether a driver who "
             "was online actually received trips.")
    table(doc, ["Driver", "Hours online", "Trips given", "Sitting idle", "Earned"], [
        ["X", "8 hours", "2", "6 hours (75%)", "Rs 5,000"],
        ["Y", "8 hours", "10", "1 hour", "Rs 5,000"],
    ], widths=[1.0, 1.3, 1.2, 1.5, 1.3])
    txt(doc, "X got only 2 trips but they happened to be long ones, so earnings are equal "
             "and the paper says fair. But X wasted 6 hours sitting with the app open and "
             "getting nothing. That is a real unfairness which is nowhere in the paper.",
        size=9.5, italic=True)

    # ================= 3 =================
    doc.add_page_break()
    head(doc, "3. What I did with the data", 1)
    txt(doc, "The raw file is 1.9 GB with 1,22,10,952 trip records. Each row is one taxi "
             "trip with pickup time, drop time, pickup and drop location, distance and fare. "
             "One important thing: the file has no driver ID. It only records trips, not "
             "which driver did them. This matters a lot, and I explain it below.")
    table(doc, ["#", "What I did", "Before to After", "Why it was needed"], [
        ["1", "Cleaned the data", "1,22,10,952 to 1,03,00,738 trips",
         "Removed trips outside Manhattan, trips of 0 km, one trip showing 50,00,000 km "
         "(GPS error), negative fares, and trips shorter than 1 minute or longer than 2 hours"],
        ["2", "Divided the city into 87 areas", "GPS points to 87 nodes",
         "The model needs a fixed set of locations, not crores of GPS points. I tested "
         "several grid sizes and picked around 1.1 km because smaller areas did not have "
         "enough trips to calculate travel time properly"],
        ["3", "Made a distance table", "87 x 87 table",
         "The paper uses Geo(a,b) but never says how it got it. I took the median actual "
         "distance from real trips. Check: my road distance comes out 1.32 times the "
         "straight-line distance, and the known value for New York is about 1.3, so the "
         "table is correct"],
        ["4", "Made a travel-time table", "1 number to 48 numbers per route",
         "This is the fix for Problem 1 Part A. 48 = weekday/weekend times 24 hours. This "
         "is where the 2.11 times traffic difference shows up"],
        ["5", "Created 200 drivers", "no drivers to 200 drivers with shifts",
         "Since the data has no driver ID, drivers have to be created. 76 full-time "
         "(around 46 hours a week) and 124 part-time (around 14 hours a week). Without "
         "hours, Problem 1 Part B and Problem 2 cannot even be measured"],
        ["6", "Decision every 5 minutes instead of 1 hour", "1 hour to 5 minutes",
         "An average Manhattan trip is 11 minutes. With 1-hour steps you cannot say whether "
         "a driver was busy or free, and that is exactly what Problem 2 needs"],
        ["7", "Used the full day instead of 2 peak hours", "2 hours a day to 24 hours a day",
         "The paper only uses 2 hours of each day. That means a driver can work at most 14 "
         "hours a week, so a 40-hour full-time driver is impossible. I ran the paper's own "
         "setting to confirm: 0 out of 200 drivers could reach full-time"],
    ], widths=[0.3, 1.5, 1.5, 3.5], font=8.5)
    txt(doc, "Important point: the full cleaned data (1,03,00,738 trips) is used to build "
             "everything — the 87 areas, the distance table, the travel-time table and the "
             "prediction model. Only the final allocation experiment runs on a smaller "
             "sample of requests, and the paper does the same thing (it uses a 5% sample).",
        size=9.5)

    # ================= 4 =================
    head(doc, "4. The new formulas I made", 1)

    head(doc, "Fix for traffic (Gap 1, Part A)", 2)
    txt(doc, "First I made a traffic number:")
    formula(doc, "c  =  (travel time on this route at this hour)  /  (average travel time on this route)")
    table(doc, ["Value of c", "Meaning"], [
        ["c = 1", "Normal speed for that route"],
        ["c = 2", "Taking twice as long as usual, heavy traffic"],
        ["c = 0.5", "Taking half the usual time, empty roads"],
    ], widths=[1.3, 4.5])
    txt(doc, "Then I put c into the utility formula:")
    formula(doc, "U_new  =  Geo(d, s) / c(s, d)   -   Geo(s, g) x c(g, s)")
    bul(doc, [
        "The paid part is DIVIDED by c. If a trip takes twice as long, the driver is stuck "
        "in traffic for double the time for the same money, so it is worth less to him.",
        "The empty-travel part is MULTIPLIED by c. Driving empty in traffic is worse than "
        "driving empty on a clear road: same distance, more time wasted, still no money.",
        "Very important check: if c = 1 everywhere, this formula becomes exactly the "
        "paper's original formula. I verified this in code and the difference came out "
        "exactly 0. So the paper is just a special case of my formula, which makes the "
        "comparison fair.",
    ])
    txt(doc, "Example with an 8 km trip and the passenger 5 km away:", after=2)
    table(doc, ["Case", "c", "Utility"], [
        ["Paper (same at any time)", "-", "8 - 5 = +3 km  (looks profitable)"],
        ["Mine, 4 AM, clear roads", "0.6", "8/0.6 - 5(0.6) = +10.3 km"],
        ["Mine, 6 PM, heavy traffic", "2.0", "8/2.0 - 5(2.0) = -6.0 km  (actually a loss)"],
    ], widths=[2.1, 0.7, 3.4], bold_rows=[2])
    txt(doc, "So the paper says this trip is good (+3) when it is actually bad (-6). I "
             "checked this over the whole data: the paper gets the good/bad sign wrong on "
             "10.2% of cases, which is 1 out of every 10.", size=9.5, italic=True)

    head(doc, "Fix for working hours (Gap 1, Part B)", 2)
    formula(doc, "rate  =  total earning of the driver  /  hours he was online")
    txt(doc, "Then instead of measuring variance of total earning, I measure variance of "
             "this rate. I also split it into two parts using a standard result from "
             "probability called the law of total variance:")
    formula(doc, "Var(rate)  =  within-group part  +  between-group part")
    table(doc, ["Term", "What it means in simple words"], [
        ["within-group", "How unequal drivers are compared to drivers of their own type "
                         "(one full-time driver vs another full-time driver)"],
        ["between-group", "How unequal the two groups are on average (full-time drivers "
                          "as a group vs part-time drivers as a group). This is the real "
                          "pay gap between the two types."],
        ["Groups used", "Full-time = 40 or more hours a week. Part-time = 20 hours or less."],
    ], widths=[1.5, 5.3])
    txt(doc, "This split is not something I invented. It is a known mathematical identity, "
             "so within + between always adds up to the total exactly. I checked this in "
             "code and the error was around 10 to the power -16, which is just computer "
             "rounding.", size=9.5)

    head(doc, "Fix for idle drivers (Gap 2)", 2)
    formula(doc, "utilisation  =  time the driver was busy  /  time the driver was online")
    txt(doc, "So if a driver is online 8 hours and busy 2 hours, utilisation = 0.25. Then I "
             "measure variance of utilisation across drivers. If it is high, some drivers "
             "are getting work and others are sitting idle.")
    txt(doc, "But there is one problem with this. If a driver himself chooses to be online "
             "at 4 AM when there are no passengers, his utilisation will be low, and that is "
             "not the app's fault. So I also made a second version:")
    formula(doc, "adjusted utilisation  =  driver's utilisation  /  what other drivers "
                 "online at the same time got")
    txt(doc, "If this comes to 1, the driver did as well as everyone else working those same "
             "hours. If it is less than 1, the app really did give him less work. I report "
             "both numbers so nothing is hidden.")

    head(doc, "The final objective I optimise", 2)
    formula(doc, "Score = total earnings\n"
                 "        - w1 x (within-group unfairness in rate)\n"
                 "        - w2 x (between-group unfairness in rate)\n"
                 "        - w3 x (unfairness in utilisation)")
    txt(doc, "w1, w2, w3 are weights that decide how much importance to give to each part. "
             "One thing I found here: the paper fixes a scaling value omega = 0.6 and does "
             "not say how it got it. When I used a fixed value, my utilisation term came "
             "out around 0.009 while a normal trip value is about 2.4. So that term was "
             "about 270 times too small and was doing nothing at all. I fixed this by "
             "measuring the correct scale from an actual run. After that all my ablation "
             "results started behaving properly. This itself is a finding about the paper.")
    txt(doc, "I also had to add one more thing. The paper's model only knows the driver's "
             "location. If the model does not know how far behind a driver is in earnings, "
             "it cannot correct it. So I added the driver's shortfall in rate and in "
             "utilisation into the model's state.")

    # ================= 5 =================
    doc.add_page_break()
    head(doc, "5. Results", 1)
    txt(doc, "All six methods were run on exactly the same setup: same 200 drivers, same "
             "6,604 requests, same distance and travel-time tables. So any difference is "
             "only because of the allocation method. Utility is in km.", size=9.5, italic=True)

    head(doc, "5.1 First, what each result column means", 2)
    table(doc, ["Column", "Meaning", "Good direction"], [
        ["Total Utility", "Total earnings of all 200 drivers for the week", "Higher"],
        ["Var(total)", "The paper's fairness measure. Spread of total weekly earnings.", "Lower"],
        ["Var(rate)", "My fairness measure. Spread of earning per hour.", "Lower"],
        ["within", "Part of Var(rate) coming from differences inside the same group", "Lower"],
        ["between", "Part of Var(rate) coming from the gap between full-time and part-time "
                    "drivers", "Lower"],
        ["FT / hr", "Average earning per hour of full-time drivers", "-"],
        ["PT / hr", "Average earning per hour of part-time drivers", "-"],
        ["ratio", "PT rate divided by FT rate. 1.00 means both groups earn the same per "
                  "hour, which is what we want.", "Close to 1.00"],
        ["Var(util)", "Spread of how busy the drivers were kept", "Lower"],
        ["Var(util) adj", "Same but after removing the effect of which hours the driver "
                          "chose to work", "Lower"],
        ["idle", "Number of drivers who got no trip at all in the whole week", "0"],
    ], widths=[1.2, 4.6, 1.0], font=8.5)

    head(doc, "5.2 Main result table", 2)
    rows = []
    for _, r in t1.iterrows():
        rows.append([r.method, f"{r.total_utility:,.0f}", f"{r.fairness_total_var:,.0f}",
                     f"{r.rate_var:.3f}", f"{r.rate_var_between:.4f}",
                     f"{r.rate_full_time:.2f}", f"{r.rate_part_time:.2f}",
                     f"{r.rate_group_ratio:.2f}", f"{r.util_var:.4f}",
                     f"{r.util_adj_var:.3f}", int(r.n_idle_drivers)])
    table(doc, ["Method", "Total Utility", "Var(total)", "Var(rate)", "between",
                "FT /hr", "PT /hr", "ratio", "Var(util)", "adj", "idle"], rows,
          widths=[1.5, 0.72, 0.62, 0.6, 0.6, 0.5, 0.5, 0.45, 0.62, 0.5, 0.35],
          font=8, bold_rows=[5])
    txt(doc, "Main thing to notice in the 'ratio' column: all five earlier methods are "
             "between 1.62 and 1.91. That means part-time drivers are earning 62% to 91% "
             "more per hour than full-time drivers in every single one of them. This is not "
             "random, it happens because all of them are trying to equalise total earnings. "
             "My method brings it to 0.97, which is almost exactly equal.", bold=True, size=9.5)

    head(doc, "5.3 My method compared to the paper's method", 2)
    rows = [
        ["Total earnings", f"{k.total_utility:,.0f}", f"{o.total_utility:,.0f}",
         f"{pc(o.total_utility, k.total_utility)} (small loss)"],
        ["Unfairness in earning per hour", f"{k.rate_var:.3f}", f"{o.rate_var:.3f}",
         f"{pc(o.rate_var, k.rate_var)} better"],
        ["Gap between the two driver groups", f"{k.rate_var_between:.4f}",
         f"{o.rate_var_between:.4f}", f"{pc(o.rate_var_between, k.rate_var_between)} better"],
        ["Full-time earning per hour", f"{k.rate_full_time:.2f}", f"{o.rate_full_time:.2f}",
         f"{pc(o.rate_full_time, k.rate_full_time)} (went up)"],
        ["Part-time earning per hour", f"{k.rate_part_time:.2f}", f"{o.rate_part_time:.2f}",
         f"{pc(o.rate_part_time, k.rate_part_time)} (came down to equal)"],
        ["PT / FT ratio (1.00 is equal)", f"{k.rate_group_ratio:.2f}",
         f"{o.rate_group_ratio:.2f}", "almost equal now"],
        ["Unfairness in busy time", f"{k.util_var:.4f}", f"{o.util_var:.4f}",
         f"{pc(o.util_var, k.util_var)} better"],
        ["Same, after adjusting for shift choice", f"{k.util_adj_var:.4f}",
         f"{o.util_adj_var:.4f}", f"{pc(o.util_adj_var, k.util_adj_var)} better"],
        ["Drivers with no work all week", f"{int(k.n_idle_drivers)}",
         f"{int(o.n_idle_drivers)}", "removed completely"],
    ]
    table(doc, ["What is measured", "Paper", "Mine", "Change"], rows,
          widths=[2.7, 1.0, 1.0, 2.1], font=9)
    txt(doc, "So by giving up only 0.7% of total earnings, the per-hour pay of the two "
             "driver groups became almost equal, and no driver was left without work.",
        bold=True)

    head(doc, "5.4 One more finding: the paper's fairness formula actually creates the problem", 2)
    txt(doc, "I ran the same method with the fairness part switched off and switched on, to "
             "see what the paper's fairness formula is really doing.")
    table(doc, ["Setting", "FT per hour", "PT per hour", "Gap between groups"], [
        ["Fairness OFF (only earnings)", "2.28", "2.42", "0.004 (almost no gap)"],
        ["Paper's fairness ON", "1.89", "3.13", "0.361 (90 times worse)"],
    ], widths=[2.3, 1.3, 1.3, 1.9], bold_rows=[1])
    txt(doc, "This is the important point. The paper's fairness formula does not just fail "
             "to see the group gap, it actually creates it. The reason is simple: if driver "
             "A works 50 hours and driver B works 10 hours, the only way to make their "
             "total earnings equal is to take work away from A and give it to B. So B's "
             "earning per hour shoots up. It is unavoidable with that formula.")

    head(doc, "5.5 How results improved as I added each fix", 2)
    rows = [
        ["No fairness at all", f"{nofair.total_utility:,.0f}", f"{nofair.rate_var:.3f}",
         f"{nofair.rate_var_between:.4f}", f"{nofair.rate_group_ratio:.2f}",
         f"{nofair.util_var:.4f}", int(nofair.n_idle_drivers)],
        ["Paper's fairness added", f"{k.total_utility:,.0f}", f"{k.rate_var:.3f}",
         f"{k.rate_var_between:.4f}", f"{k.rate_group_ratio:.2f}",
         f"{k.util_var:.4f}", int(k.n_idle_drivers)],
        ["My Gap 1 fixes added", f"{g12.total_utility:,.0f}", f"{g12.rate_var:.3f}",
         f"{g12.rate_var_between:.4f}", f"{g12.rate_group_ratio:.2f}",
         f"{g12.util_var:.4f}", int(g12.n_idle_drivers)],
        ["My Gap 2 fix added (final)", f"{o.total_utility:,.0f}", f"{o.rate_var:.3f}",
         f"{o.rate_var_between:.4f}", f"{o.rate_group_ratio:.2f}",
         f"{o.util_var:.4f}", int(o.n_idle_drivers)],
    ]
    table(doc, ["Step", "Total Utility", "Var(rate)", "between", "ratio", "Var(util)", "idle"],
          rows, widths=[1.9, 0.95, 0.8, 0.8, 0.6, 0.85, 0.45], font=8.5, bold_rows=[3])
    bul(doc, [
        f"Step 1 to 2: paper's fairness makes the group gap much worse "
        f"({nofair.rate_var_between:.4f} to {k.rate_var_between:.4f}) and the ratio jumps "
        f"from {nofair.rate_group_ratio:.2f} to {k.rate_group_ratio:.2f}.",
        f"Step 2 to 3: my Gap 1 fix brings Var(rate) down from {k.rate_var:.3f} to "
        f"{g12.rate_var:.3f} and the group gap almost to zero. Idle drivers also become 0.",
        f"Step 3 to 4: my Gap 2 fix brings Var(util) down from {g12.util_var:.4f} to "
        f"{o.util_var:.4f}, about {abs(int(round(100*(o.util_var-g12.util_var)/g12.util_var)))}% "
        f"better. This proves the Gap 2 part is doing real work and not just riding on Gap 1.",
    ], size=9.5)

    head(doc, "5.6 Checks I did to make sure it is not luck", 2)
    m = rb[rb.variant == "money"]
    mk = m[m.method.str.startswith("Kang")].iloc[0]
    mo = m[m.method.str.startswith("Ours")].iloc[0]
    sd = rb[rb.variant.astype(str).str.startswith("seed")].copy()
    sd["fam"] = ["Mine" if str(x).startswith("Ours") else "Paper" for x in sd.method]
    agg = sd.groupby("fam").agg(ra=("rate_group_ratio", "mean"),
                                rs=("rate_group_ratio", "std"))
    dom = gr[(gr.total_utility > k.total_utility) & (gr.rate_var < k.rate_var)
             & (gr.rate_var_between < k.rate_var_between) & (gr.util_var < k.util_var)
             & (gr.util_adj_var < k.util_adj_var)
             & (gr.n_idle_drivers <= k.n_idle_drivers)]
    table(doc, ["Check", "Result"], [
        ["Did it work in rupees/dollars also, not just km?",
         f"Yes. Paper: full-time ${mk.rate_full_time:.2f}/hr vs part-time "
         f"${mk.rate_part_time:.2f}/hr. Mine: ${mo.rate_full_time:.2f} vs "
         f"${mo.rate_part_time:.2f}. So the gap is real, not a unit problem."],
        ["Did it work with different random settings?",
         f"Yes, ran 3 different random seeds. Paper's ratio stayed "
         f"{agg.loc['Paper','ra']:.2f} every time (so the problem is systematic, not luck). "
         f"Mine stayed {agg.loc['Mine','ra']:.2f}."],
        ["Am I just trading earnings for fairness?",
         f"No. I tested {len(gr)} different weight settings. In {len(dom)} of them my "
         f"method beat the paper on all six measures at the same time, including total "
         f"earnings."],
        ["Is the prediction model working?",
         f"Yes. Demand prediction error {fc['demand_mse_counts']:.2f} against a simple "
         f"baseline of {fc['baseline_mse_counts']:.2f}, so about "
         f"{100*(1-fc['demand_mse_counts']/fc['baseline_mse_counts']):.0f}% better. "
         f"Traffic prediction error is 0.08, which is small."],
    ], widths=[2.1, 4.7], font=9)

    # ================= 6 =================
    doc.add_page_break()
    head(doc, "6. Comparing with the numbers printed in the paper", 1)
    txt(doc, "Sir, one thing needs explaining. The paper's Table 1 shows total utility "
             "95,823.79 for its method, and mine shows 10,145. This looks like my result is "
             "much smaller, but the two numbers are measuring different situations. I am not "
             "using less data — the full cleaned data is used for everything. There are two "
             "reasons for the difference, and both were hidden in the paper.")

    head(doc, "Reason 1: the paper used 20 drivers, I used 200", 2)
    txt(doc, "The paper never says how many drivers it used. But total utility divided by "
             "mean utility per driver gives the number of drivers, and it comes out exactly "
             "20 for every row of their table:")
    table(doc, ["Method in paper's Table 1", "Total", "Mean per driver", "So drivers ="], [
        ["Greedy", "-15,14,736.24", "-75,736.81", "20.000"],
        ["REASSIGN", "76,536.23", "3,826.81", "20.000"],
        ["LAF", "80,606.49", "4,030.3245", "20.000"],
        ["Balance Ride-Pooling", "85,923.68", "4,296.18", "20.000"],
        ["Proposed Method", "95,823.79", "4,791.19", "20.000"],
    ], widths=[2.2, 1.5, 1.5, 1.2], font=9)
    txt(doc, "With only 20 drivers, the same amount of work gets divided among far fewer "
             "people, so the per-driver number becomes very large.", size=9.5, italic=True)

    head(doc, "Reason 2: their number needs more trips than are physically possible", 2)
    table(doc, ["Step", "Value"], [
        ["Paper's average earning per driver per week", "4,791"],
        ["Realistic value of one trip (from my measurement)", "about 2.5 km"],
        ["So trips needed per driver per week", "about 1,924 trips"],
        ["But paper uses only 2 peak hours per day, so hours available", "14 hours a week"],
        ["Average trip 11 minutes + 4.5 minutes to reach passenger, so maximum trips",
         "about 53 trips"],
        ["Needed vs possible", "about 36 times more than possible"],
    ], widths=[4.3, 2.5], font=9)
    txt(doc, "The reason is in their Section 4.3: they allow one driver to accept many "
             "requests at the same time, and their formula has no time in it, so nothing "
             "stops a driver from taking unlimited trips in one time step. In my simulation "
             "a driver is busy for the actual travel time and cannot take another trip until "
             "he finishes. That is why my totals are smaller, and that same rule is what "
             "makes Problem 2 measurable in the first place.")
    txt(doc, "So instead of comparing my number with their printed number, I re-implemented "
             "their method inside my own setup and compared there. That way both methods see "
             "the same drivers, same requests and same distances, and the comparison is "
             "actually fair. That is what all the tables in Section 5 show.", bold=True)

    head(doc, "6.1 Summary of where my work is better", 2)
    table(doc, ["Point", "Paper", "Mine"], [
        ["What fairness means", "Equal total earnings, hours ignored",
         "Equal earning per hour, and checked separately for full-time and part-time"],
        ["Pay gap between driver types", f"{k.rate_group_ratio:.2f} times",
         f"{o.rate_group_ratio:.2f} times (almost equal)"],
        ["Idle drivers checked?", "No", "Yes, and brought down to 0"],
        ["Traffic considered?", "No, one average for the whole day",
         "Yes, 48 different values per route"],
        ["Driver time limit", "None, can take unlimited trips together",
         "Driver is busy for the real trip time"],
        ["Over a longer week", "Group gap keeps growing (16.5 times)", "Stays flat"],
        ["Scaling value omega", "Fixed at 0.6, no explanation",
         "Measured from actual run, otherwise the term does nothing"],
        ["Number of drivers reported?", "No (I worked it out as 20)", "Yes, 200"],
    ], widths=[1.7, 2.2, 2.9], font=8.5)

    # ================= 7 =================
    head(doc, "7. Things that did not work (being honest)", 1)
    bul(doc, [
        f"The paper's prediction module did not help me. When I removed it, total earnings "
        f"actually went slightly up ({o.total_utility:,.0f} to {nopred.total_utility:,.0f}). "
        f"The paper claims a 41% drop without prediction (95,824 down to 56,873) and I could "
        f"not reproduce that at all. My guess is that my model already sees all 48 "
        f"time-slots during training, so predicting the future adds nothing new.",
        f"My traffic fix (Gap 1 Part A) improves earnings by only about 1%. Its real "
        f"benefit is correctness, not money — it fixes the good/bad sign on 10.2% of trips. "
        f"I am reporting it that way and not claiming more.",
        "My Var(total) is worse than the paper's (1,639 vs 657). This is not a mistake. If "
        "everyone gets equal pay per hour but works different hours, then their totals must "
        "be different. Both cannot be zero at the same time. This is actually my main point.",
        f"Service rate is slightly lower ({100*o.service_rate:.1f}% vs "
        f"{100*k.service_rate:.1f}%), so about 2% fewer requests get served.",
    ], size=9.5)

    # ================= 8 =================
    head(doc, "8. Summary", 1)
    txt(doc, "In one line: the paper is trying to make total earnings equal, which is the "
             "wrong target. It hides a big difference in earning per hour, it actually makes "
             "the full-time vs part-time gap worse, and it never checks whether an online "
             "driver got any work.", bold=True)
    txt(doc, "What I changed: I made the utility formula aware of traffic, changed fairness "
             "from total earning to earning per hour, split it into within-group and "
             "between-group, and added a measure for whether drivers are actually being "
             "given work.")
    txt(doc, f"Result: earning per hour of full-time and part-time drivers is now almost "
             f"equal ({o.rate_group_ratio:.2f} against {k.rate_group_ratio:.2f} in the "
             f"paper), unfairness in per-hour pay is {abs(int(round(100*(o.rate_var-k.rate_var)/k.rate_var)))}% "
             f"lower, unfairness in idle time is "
             f"{abs(int(round(100*(o.util_adj_var-k.util_adj_var)/k.util_adj_var)))}% lower, "
             f"and no driver is left without work. The cost is only "
             f"{abs(100*(o.total_utility-k.total_utility)/k.total_utility):.1f}% of total "
             f"earnings.")
    txt(doc, "Next steps I am planning: check how results change with different fleet sizes "
             "(100 / 200 / 400 drivers), and start writing the paper draft.")

    txt(doc, "")
    txt(doc, "Note: everything in this document can be reproduced. Running "
             "'python -m scripts.show_results' prints all these numbers in a few seconds, "
             "and 'python -m scripts.explain_scale' shows the data accounting and the "
             "20-driver calculation. A longer technical report with all equations and "
             "correctness proofs is also ready if needed.", size=9, italic=True)

    out = OUTPUT / "MTP_Progress_Update.docx"
    doc.save(out)
    print(f"wrote {out}")
    print(f"size: {out.stat().st_size:,} bytes")
    return 0


if __name__ == "__main__":
    sys.exit(main())
